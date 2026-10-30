"""
Testes do fluxo completo, com os números reais de Vila Ramos (setembro de 2026).
Rodar: python -m pytest -q
Com Postgres: DATABASE_URL=postgresql://... python -m pytest -q
"""
import io
from datetime import date
import os
import re
import sys
import tempfile

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

if not os.environ.get('DATABASE_URL'):
    os.environ['SQLITE_PATH'] = os.path.join(tempfile.mkdtemp(), 'teste.db')
os.environ['ADMIN_USUARIO'] = 'nacional'
os.environ['ADMIN_SENHA'] = 'senha-forte-123'
os.environ['COOKIE_SECURE'] = '0'

import app as appmod  # noqa: E402
import regras as R  # noqa: E402

PDF = b'%PDF-1.4\n% comprovante de teste\n'


@pytest.fixture(scope='module')
def app():
    appmod.app.config['TESTING'] = True
    appmod.app.config['TESTING_SEM_CSRF'] = True
    return appmod.app


def entrar(app, usuario, senha):
    c = app.test_client()
    r = c.post('/login', data={'usuario': usuario, 'senha': senha})
    assert r.status_code == 302, r.data[:300]
    return c


# ── Motor de regras ─────────────────────────────────────────
def test_carta_vila_ramos_2026():
    meses = [6774.70, 5522.40, 6315.00, 4477.00, 3181.50, 5218.40, 3450.60, 3450.00, 5204.70, 7361.40, 8650.00, 0]
    reg = {'pct_dizimo_dizimos': .10, 'fator_formacao': 1 / 3, 'fator_nivelacao': 1.25}
    c = R.carta_anual(meses, 1518, reg)
    assert c['media'] == 5418.70
    assert c['categoria'] == 'formada'
    assert c['fundo_local_pct'] == 3
    assert c['nivelacao'] is True


def test_categorias_pela_entrada():
    reg = {'pct_dizimo_dizimos': .10, 'fator_formacao': 1 / 3, 'fator_nivelacao': 1.25}
    assert R.carta_anual([0] * 12, 1621, reg)['categoria'] == 'nova'
    assert R.carta_anual([0] * 12, 1621, reg)['fundo_local_pct'] == 0
    assert R.carta_anual([500], 1621, reg)['categoria'] == 'projecao'
    assert R.carta_anual([600], 1621, reg)['categoria'] == 'formacao'
    assert R.carta_anual([1621], 1621, reg)['categoria'] == 'formada'


def test_fechamento_vila_ramos_setembro():
    reg = {'pct_dizimo_dizimos': .10, 'pct_nivelacao': .05, 'pct_betel': .04, 'min_invest_formada': 500}
    comp = {'categoria': 'formada', 'nivelacao': True, 'fundo_local_pct': 3, 'betel': True}
    envios = {'Of. Dia Missionário': 817.80, 'Of. Raios De Luz': 457.60, 'Of. Safras Para Cristo': 115.20,
              'Of. Distrital': 150.00, 'Voto Pastor': 50.00}
    calc = R.fechar_mes(7736.47, comp, envios, reg)
    liq = calc['liquidacao']
    assert (liq['dd'], liq['niv'], liq['betel'], liq['fundo_local']) == (773.65, 386.82, 263.04, 232.09)
    assert liq['liquido'] == 6080.87
    assert calc['total'] == 3014.11   # igual ao relatório e ao comprovante reais
    assert any('Fundo de investimento' in a for a in calc['avisos'])


def test_prazo():
    assert str(R.prazo_envio(2026, 9)) == '2026-10-05'
    assert str(R.prazo_envio(2026, 12)) == '2027-01-05'


# ── Fluxo completo no sistema ───────────────────────────────
def test_fluxo_completo(app):
    adm = entrar(app, 'nacional', 'senha-forte-123')
    r = adm.post('/admin/nova', data={'nome': 'Vila Ramos/SP',
                                     'tes_nome': 'José Antonio Cervera', 'tes_usuario': 'vr.tesouraria', 'tes_senha': 'tesoureiro123',
                                     'pas_nome': 'Carlos Javier Hernández', 'pas_usuario': 'vr.pastor', 'pas_senha': 'pastor12345'})
    assert r.status_code == 302
    with appmod.get_db() as conn:
        cid = conn.execute("SELECT id FROM congregacoes WHERE nome='Vila Ramos/SP'").fetchone()['id']
        # dízimos de 2025 (carta real) para o sistema calcular a carta de 2026
        for i, v in enumerate([6774.70, 5522.40, 6315.00, 4477.00, 3181.50, 5218.40, 3450.60, 3450.00, 5204.70, 7361.40, 8650.00]):
            conn.execute('INSERT INTO dizimos (cong_id, dizimo_bruto, mes, ano) VALUES (?,?,?,?)', (cid, v, R.MESES[i], 2025))
    # salário mínimo de 2025 é a base da carta de 2026
    r = adm.post('/nacional/cartas', data={'ano': 2026, 'acao': 'emitir', 'cong_id': cid})
    assert r.status_code == 302
    with appmod.get_db() as conn:
        c = conn.execute('SELECT * FROM compromissos WHERE cong_id=? AND ano=2026', (cid,)).fetchone()
    assert (c['categoria'], c['fundo_local_pct'], c['nivelacao'], c['betel']) == ('formada', 3, 1, 1)

    pas = entrar(app, 'vr.pastor', 'pastor12345')
    assert pas.post('/carta/2026/confirmar').status_code == 302

    tes = entrar(app, 'vr.tesouraria', 'tesoureiro123')
    # o pastor não lança
    assert pas.post('/dizimos', data={'dizimo_bruto': '100', 'mes': 'setembro', 'ano': 2026}).status_code == 403
    tes.post('/dizimos', data={'dizimo_bruto': '7.736,47', 'mes': 'setembro', 'ano': 2026})
    for item, valor in [('Of. Dia Missionário', '817,80'), ('Of. Raios De Luz', '457,60'),
                        ('Of. Safras Para Cristo', '115,20'), ('Of. Distrital', '150,00'), ('Voto Pastor', '50,00')]:
        r = tes.post('/lancamentos', data={'livro': 'Envios', 'tipo': 'Entrada', 'item': item, 'valor': valor,
                                           'mes': 'setembro', 'ano': 2026})
        assert r.status_code == 302
    # uma despesa local com comprovante
    tes.post('/lancamentos', data={'livro': 'FundoLocal', 'tipo': 'Saída', 'item': 'Serviços Públicos (água, luz)',
                                   'valor': '210,00', 'mes': 'setembro', 'ano': 2026,
                                   'anexo': (io.BytesIO(PDF), 'conta-luz.pdf', 'application/pdf')},
             content_type='multipart/form-data')

    pagina = tes.get('/fechamento?ano=2026&mes=9').get_data(as_text=True)
    assert 'R$ 3.014,11' in pagina

    assert tes.post('/fechamento/fechar', data={'ano': 2026, 'mes': 9}).status_code == 302
    with appmod.get_db() as conn:
        f = conn.execute('SELECT * FROM fechamentos WHERE cong_id=? AND ano=2026 AND mes=9', (cid,)).fetchone()
    assert f['status'] == 'fechado' and f['total'] == 3014.11

    # mês travado: não aceita lançamento nem dízimo
    tes.post('/lancamentos', data={'livro': 'FundoLocal', 'tipo': 'Entrada', 'item': 'Outras Ofertas',
                                   'valor': '10', 'mes': 'setembro', 'ano': 2026})
    with appmod.get_db() as conn:
        n = conn.execute("SELECT COUNT(*) AS n FROM transacoes WHERE cong_id=? AND item='Outras Ofertas'", (cid,)).fetchone()['n']
    assert n == 0

    # tesoureiro não aprova; pastor aprova
    assert tes.post(f"/fechamento/{f['id']}/aprovar").status_code == 302
    with appmod.get_db() as conn:
        assert conn.execute('SELECT status FROM fechamentos WHERE id=?', (f['id'],)).fetchone()['status'] == 'fechado'
    pas.post(f"/fechamento/{f['id']}/aprovar")

    # comprovante errado é recusado
    tes.post(f"/fechamento/{f['id']}/comprovante", data={'valor': '3.000,00', 'data': '2026-10-07',
             'arquivo': (io.BytesIO(PDF), 'tev.pdf', 'application/pdf')}, content_type='multipart/form-data')
    with appmod.get_db() as conn:
        assert conn.execute('SELECT status FROM fechamentos WHERE id=?', (f['id'],)).fetchone()['status'] == 'aprovado'
    tes.post(f"/fechamento/{f['id']}/comprovante", data={'valor': '3.014,11', 'data': '2026-10-07',
             'arquivo': (io.BytesIO(PDF), 'tev.pdf', 'application/pdf')}, content_type='multipart/form-data')
    with appmod.get_db() as conn:
        assert conn.execute('SELECT status FROM fechamentos WHERE id=?', (f['id'],)).fetchone()['status'] == 'enviado'
        gerados = conn.execute('SELECT * FROM transacoes WHERE origem=?', (f"fechamento:{f['id']}",)).fetchall()
    itens = {g['item']: g['valor'] for g in gerados}
    assert itens['Envio Dízimo dos Dízimos 10%'] == 773.65
    assert itens['Aporte ao fundo local'] == 232.09
    assert itens['Sustento pastoral (§ 66)'] == 6080.87

    # relatório econômico e comprovante
    rel = tes.get(f"/fechamento/{f['id']}/relatorio").get_data(as_text=True)
    assert 'R$ 3.014,11' in rel and 'base R$ 6.576,00' in rel and 'logo-inj.png' in rel
    assert tes.get(f"/fechamento/{f['id']}/comprovante.arquivo").data == PDF

    # visão nacional e confirmação de recebimento
    painel = adm.get('/nacional/envios?ano=2026').get_data(as_text=True)
    assert 'Vila Ramos/SP' in painel
    adm.post(f"/nacional/fechamento/{f['id']}/receber")
    with appmod.get_db() as conn:
        assert conn.execute('SELECT status FROM fechamentos WHERE id=?', (f['id'],)).fetchone()['status'] == 'recebido'

    # depois de recebido, o pastor não reabre
    pas.post(f"/fechamento/{f['id']}/reabrir", data={'motivo': 'teste de reabertura'})
    with appmod.get_db() as conn:
        assert conn.execute('SELECT 1 FROM fechamentos WHERE id=?', (f['id'],)).fetchone()

    # outro mês: fechar e reabrir com motivo apaga os lançamentos gerados
    tes.post('/dizimos', data={'dizimo_bruto': '5000', 'mes': 'outubro', 'ano': 2026})
    tes.post('/fechamento/fechar', data={'ano': 2026, 'mes': 10})
    with appmod.get_db() as conn:
        f10 = conn.execute('SELECT id FROM fechamentos WHERE cong_id=? AND ano=2026 AND mes=10', (cid,)).fetchone()
    pas.post(f"/fechamento/{f10['id']}/reabrir", data={'motivo': 'faltou lançar a oferta missionária'})
    with appmod.get_db() as conn:
        assert not conn.execute('SELECT 1 FROM fechamentos WHERE id=?', (f10['id'],)).fetchone()
        aud = conn.execute("SELECT detalhe FROM auditoria WHERE acao='fechamento.reabrir'").fetchone()
    assert 'faltou lançar' in aud['detalhe']

    # todas as páginas abrem para cada perfil
    for cli, paginas in [
        (tes, ['/', '/lancamentos', '/dizimos', '/fechamento', '/carta?ano=2026', '/orcamento', '/relatorio',
               '/configuracoes', '/minha-conta', '/tutorial', '/exportar/excel?ano=2026', '/lancar', '/lancar?tipo=Entrada',
               '/lancar?tipo=Saída', '/visao-geral', '/lancamentos?livro=Projetos', '/livro-do-mes?ano=2026&mes=9']),
        (pas, ['/', '/lancamentos', '/dizimos', '/fechamento?ano=2026&mes=9', '/carta?ano=2026']),
        (adm, ['/nacional/envios', '/nacional/cartas?ano=2027', '/nacional/regras', '/admin', '/nacional/auditoria',
               f'/nacional/carta/{cid}/2026']),
    ]:
        for p in paginas:
            r = cli.get(p, follow_redirects=True)
            assert r.status_code == 200, (p, r.status_code)

    # congregação não vê o painel nacional
    r = tes.get('/nacional/envios')
    assert r.status_code == 302


def test_seguranca(app):
    c = app.test_client()
    # sem senha fixa no código
    assert 'admin123' not in open(os.path.join(ROOT, 'app.py'), encoding='utf-8').read()
    # bloqueio depois de 5 tentativas erradas
    for _ in range(5):
        c.post('/login', data={'usuario': 'nacional', 'senha': 'errada'})
    r = c.post('/login', data={'usuario': 'nacional', 'senha': 'senha-forte-123'})
    assert r.status_code == 200 and 'Muitas tentativas' in r.get_data(as_text=True)
    # formulário sem token CSRF é recusado
    app.config['TESTING_SEM_CSRF'] = False
    try:
        c2 = app.test_client()
        r = c2.post('/login', data={'usuario': 'vr.tesouraria', 'senha': 'tesoureiro123'})
        assert r.status_code == 302 and '/inicio' not in r.headers['Location']
        pagina = c2.get('/login').get_data(as_text=True)
        token = re.search(r'name="csrf_token" value="([^"]+)"', pagina).group(1)
        r = c2.post('/login', data={'usuario': 'vr.tesouraria', 'senha': 'tesoureiro123', 'csrf_token': token})
        assert r.status_code == 302 and r.headers['Location'].endswith('/inicio')
    finally:
        app.config['TESTING_SEM_CSRF'] = True


def test_foto_de_comprovante_e_reduzida(app):
    from PIL import Image
    big = io.BytesIO()
    Image.new('RGB', (4000, 3000), (200, 30, 30)).save(big, 'PNG')
    with app.test_request_context('/', method='POST', data={'anexo': (io.BytesIO(big.getvalue()), 'foto.png', 'image/png')},
                                  content_type='multipart/form-data'):
        a = appmod._ler_anexo('anexo')
    assert a['tipo'] == 'image/jpeg' and a['nome'] == 'foto.jpg'
    assert len(a['dados']) < 300 * 1024
    assert Image.open(io.BytesIO(a['dados'])).size[0] <= 1600


def test_inicio_guiado_e_livro_do_mes(app):
    tes = entrar(app, 'vr.tesouraria', 'tesoureiro123')
    inicio = tes.get('/?ano=2026&mes=9').get_data(as_text=True)
    assert 'Os lançamentos começam aqui' in inicio and '>ENTRADA<' in inicio and '>SAÍDA<' in inicio and 'Orçamento do ano' in inicio
    r = tes.post('/lancamentos', data={'livro': 'Projetos', 'tipo': 'Entrada', 'item': 'Brechó', 'valor': '85,50',
                                       'mes': 'novembro', 'ano': 2026, 'data_registro': '02/11/2026', 'voltar': 'lancar'})
    assert '/lancar' in r.headers['Location']
    livro = tes.get('/livro-do-mes?ano=2026&mes=9').get_data(as_text=True)
    assert 'LIVRO CAIXA DO MÊS' in livro and 'Serviços Públicos' in livro and 'Saldo anterior' in livro
    assert 'Brechó' in tes.get('/livro-do-mes?ano=2026&mes=11').get_data(as_text=True)


def test_admin_abre_congregacao(app):
    with appmod.get_db() as conn:
        conn.execute('DELETE FROM login_falhas')
    adm = entrar(app, 'nacional', 'senha-forte-123')
    with appmod.get_db() as conn:
        cid = conn.execute("SELECT id FROM congregacoes WHERE nome='Vila Ramos/SP'").fetchone()['id']
    assert adm.get(f'/ver/{cid}').status_code == 302
    pagina = adm.get('/').get_data(as_text=True)
    assert 'Vila Ramos/SP' in pagina and '>ENTRADA<' in pagina
    assert adm.get('/lancar?tipo=Saída').status_code == 200
    adm.get('/ver/nacional')
    assert adm.get('/').status_code == 302


def test_recibo_configuracao_balanco_e_impressao(app):
    with appmod.get_db() as conn:
        conn.execute('DELETE FROM login_falhas')
    tes = entrar(app, 'vr.tesouraria', 'tesoureiro123')
    # nomes dos responsáveis
    r = tes.post('/configuracoes', data={'acao': 'dados', 'nome_congregacao': 'Vila Ramos/SP', 'nome_pastor': 'Pr. Fulano de Tal',
                                         'nome_tesoureiro': 'Irmã Beltrana', 'ano_atual': 2026, 'saldo_ano_anterior': '1.000,00'})
    assert r.status_code == 302
    cfg = tes.get('/configuracoes?ano=2026').get_data(as_text=True)
    assert 'Pr. Fulano de Tal' in cfg and 'Confirmada pelo pastor' in cfg  # carta 2026 já confirmada: travada p/ tesoureiro
    assert 'porcentagens' in tes.post('/configuracoes', data={'acao': 'porcentagens', 'ano': 2026, 'categoria': 'nova'},
                                      follow_redirects=True).get_data(as_text=True).lower()
    with appmod.get_db() as conn:
        assert conn.execute("SELECT categoria FROM compromissos WHERE ano=2026").fetchone()['categoria'] == 'formada'
    # ano sem carta: congregação informa, pastor confirma
    # ano sem carta: com os dízimos do ano anterior o sistema define classificação e liquidação
    dados = {'acao': 'porcentagens', 'ano': 2027}
    dados.update({f'm{i}': '900,00' for i in range(12)})   # média 900 (SM 2026 = 1.621): em formação, sem Betel
    tes.post('/configuracoes', data=dados)
    with appmod.get_db() as conn:
        c = conn.execute("SELECT * FROM compromissos WHERE ano=2027").fetchone()
        assert (c['categoria'], c['fundo_local_pct'], c['betel'], c['nivelacao'], c['confirmado_em'], c['meses']) == ('formacao', 0, 0, 0, None, 12)
    pas = entrar(app, 'vr.pastor', 'pastor12345')
    dados.update({f'm{i}': '2.500,00' for i in range(12)})  # 2.500 > 1,25 SM: formada, nivelação, Betel, fundo 1%
    pas.post('/configuracoes', data=dados)
    with appmod.get_db() as conn:
        c = conn.execute("SELECT * FROM compromissos WHERE ano=2027").fetchone()
        assert (c['categoria'], c['fundo_local_pct'], c['betel'], c['nivelacao']) == ('formada', 1, 1, 1) and c['confirmado_em']
    assert 'já foram confirmadas' in tes.post('/configuracoes', data=dados, follow_redirects=True).get_data(as_text=True)
    # lançamento com nome -> recibo com valor por extenso e nomes
    r = tes.post('/lancamentos', data={'livro': 'FundoLocal', 'tipo': 'Saída', 'item': 'Serviços Públicos (água, luz)', 'valor': '1.250,50',
                                       'mes': 'outubro', 'ano': 2026, 'pessoa': 'Sabesp', 'voltar': 'lancar'})
    tid = int(r.headers['Location'].split('ultimo=')[1])
    assert 'Imprimir recibo' in tes.get(r.headers['Location']).get_data(as_text=True)
    rec = tes.get(f'/recibo/lancamento/{tid}').get_data(as_text=True)
    assert 'mil duzentos e cinquenta reais e cinquenta centavos' in rec and 'Sabesp' in rec and 'Irmã Beltrana' in rec
    with appmod.get_db() as conn:
        did = conn.execute('SELECT id FROM dizimos ORDER BY id LIMIT 1').fetchone()['id']
    assert 'Dízimos' in tes.get(f'/recibo/dizimo/{did}').get_data(as_text=True)
    assert tes.get('/recibo/lancamento/999999').status_code == 404
    # livros, balanço e central de impressão
    ano = tes.get('/livro-do-mes?ano=2026&periodo=ano').get_data(as_text=True)
    assert 'LIVRO CAIXA DO ANO' in ano and 'Pr. Fulano de Tal' in ano and 'Serviços Públicos' in ano
    so = tes.get('/livro-do-mes?ano=2026&mes=10&livro=Projetos').get_data(as_text=True)
    assert 'LIVRO PROJETOS' in so and 'Sabesp' not in so
    bal = tes.get('/balanco?ano=2026').get_data(as_text=True)
    assert 'Balanço por livro' in bal and 'g-mensal' in bal and 'Sustento pastoral' in bal
    assert tes.get('/balanco?ano=2026&mes=9').status_code == 200
    assert 'Relatório dos envios' in tes.get('/imprimir?ano=2026&mes=9').get_data(as_text=True)


def test_valor_extenso():
    assert appmod.valor_extenso(3014.11) == 'três mil e quatorze reais e onze centavos'
    assert appmod.valor_extenso(100) == 'cem reais'
    assert appmod.valor_extenso(1) == 'um real'


def test_cadastro_obrigatorio_livro_excel_e_tutorial(app):
    import openpyxl
    with appmod.get_db() as conn:
        conn.execute('DELETE FROM login_falhas')
    adm = entrar(app, 'nacional', 'senha-forte-123')
    adm.post('/admin/nova', data={'nome': 'Nova Esperança', 'tes_nome': 'Maria Souza', 'tes_usuario': 'ne.tes', 'tes_senha': 'tesoureiro123',
                                  'pas_nome': 'Pr. João', 'pas_usuario': 'ne.pas', 'pas_senha': 'pastor12345'})
    tes = entrar(app, 'ne.tes', 'tesoureiro123')
    # nomes já vêm da criação; sem a classificação do ano, nada é liberado
    r = tes.get('/lancar?tipo=Entrada')
    assert r.status_code == 302 and '/configuracoes' in r.headers['Location']
    r = tes.post('/lancamentos', data={'livro': 'FundoLocal', 'tipo': 'Entrada', 'item': 'Votos', 'valor': '10', 'mes': 'outubro', 'ano': date.today().year})
    assert '/configuracoes' in r.headers['Location']
    assert '/configuracoes' in tes.get('/dizimos').headers['Location']
    pagina = tes.get('/configuracoes').get_data(as_text=True)
    assert 'Maria Souza' in pagina and 'Pr. João' in pagina
    # congregação nova: tudo em zero, começa do zero
    tes.post('/configuracoes', data={'acao': 'porcentagens', 'ano': date.today().year})
    with appmod.get_db() as conn:
        c = conn.execute("SELECT c.* FROM compromissos c JOIN congregacoes g ON g.id=c.cong_id WHERE g.nome='Nova Esperança'").fetchone()
        assert c['categoria'] == 'nova' and c['fundo_local_pct'] == 0
    pag = tes.get('/lancar?tipo=Saída').get_data(as_text=True)
    assert 'Em qual <b>livro</b>' in pag and 'livro-escolha' in pag
    # exportação
    r = tes.get(f'/exportar/excel?ano={date.today().year}')
    wb = openpyxl.load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames[:3] == ['Cadastro', 'Lançamentos', 'Caixa local'] and 'Dízimos' in wb.sheetnames and 'Balanço' in wb.sheetnames
    assert 'Ordem do trabalho'.lower() in tes.get('/tutorial').get_data(as_text=True).lower()


def test_itens_completos_e_nomes_antigos(app):
    # itens da planilha de exemplo e do sistema antigo
    for item in ('Conta de Água', 'Conta de Luz', 'Tarifas Bancárias', 'IPTU da Igreja', 'Visitas Pastorais',
                 'Poupança para Assembleia de Pastores', 'Transferência para o Fundo de Construção'):
        assert item in appmod.CATEGORIAS['FundoLocal']['Saída']
    assert 'Oferta das Crianças' in appmod.CATEGORIAS['FundoLocal']['Entrada']
    for item in ('Of. Raios De Luz', 'Of. Safras Para Cristo', 'Of. Dia Missionário', 'Of. Damas Miss. Metade Ano',
                 'Of. Jovens Miss. Fim Ano', 'Of. Distrital', 'Sócio Missionário', 'Fundo de Investimento'):
        assert item in appmod.CATEGORIAS['Envios']['Entrada']
    saidas_orc = [i for _, its in appmod.ORC_GRUPOS['Saída'] for i in its]
    assert 'Envio: Raios de Luz' in saidas_orc and 'Sustento pastoral (§ 66)' in saidas_orc
    # nome antigo vira o atual
    with appmod.get_db() as conn:
        conn.execute("INSERT INTO transacoes (cong_id,livro,tipo,item,valor,mes,ano) VALUES (1,'FundoLocal','Saída','Aluguel Salão',10,'janeiro',2026)")
    appmod.init_db()
    with appmod.get_db() as conn:
        assert not conn.execute("SELECT 1 FROM transacoes WHERE item='Aluguel Salão'").fetchone()
    with appmod.get_db() as conn:
        conn.execute('DELETE FROM login_falhas')
    tes = entrar(app, 'vr.tesouraria', 'tesoureiro123')
    pag = tes.get('/lancar?tipo=Saída').get_data(as_text=True)
    assert 'saídas do livro Envios' in pag and 'Raios de Luz' in pag
    orc = tes.get('/orcamento?ano=2026').get_data(as_text=True)
    assert 'Liquidação dos dízimos' in orc and 'Envio: Safras para Cristo' in orc
